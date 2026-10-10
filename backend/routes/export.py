import csv
import io
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import JSONResponse
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pymongo.database import Database

from backend.routes.dependencies import get_auth_context, get_database


router = APIRouter(prefix="/api/export", tags=["export"])


def _serialize_doc(doc: dict) -> dict:
    res = {}
    for k, v in doc.items():
        if k in ("_id", "pin_hash"):
            if k == "_id":
                res["id"] = str(v)
            continue
        if isinstance(v, datetime):
            res[k] = v.isoformat()
        elif isinstance(v, dict):
            res[k] = _serialize_doc(v)
        elif isinstance(v, list):
            res[k] = [(_serialize_doc(item) if isinstance(item, dict) else item) for item in v]
        else:
            res[k] = v
    return res


def _format_dt(dt: datetime | str | None) -> str:
    if not dt:
        return ""
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%d %I:%M %p")
    try:
        parsed = datetime.fromisoformat(str(dt).replace("Z", "+00:00"))
        return parsed.strftime("%Y-%m-%d %I:%M %p")
    except Exception:
        return str(dt)


def _build_excel_export(user: dict, logs: list, days: list, saved_meals: list) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10, color="0F172A")
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")
    data_align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    def apply_sheet_styling(ws, header_cols: list[str]):
        ws.row_dimensions[1].height = 26
        for col_num, title in enumerate(header_cols, 1):
            cell = ws.cell(row=1, column=col_num, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ── SHEET 1: Daily Logs ───────────────────────────────────────────────────
    ws_logs = wb.create_sheet(title="Daily Logs")
    logs_headers = [
        "Date",
        "Time Logged",
        "Meal / Food Entry",
        "Calories (kcal)",
        "Protein (g)",
        "Carbs (g)",
        "Fats (g)",
        "Water (ml)",
        "Workout Activity",
        "Source",
    ]
    ws_logs.row_dimensions[1].height = 26
    for c_idx, h in enumerate(logs_headers, 1):
        cell = ws_logs.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    row_num = 2
    for log in logs:
        extracted = log.get("extracted") or {}
        nutrition = extracted.get("nutrition") or {}
        overrides = log.get("overrides") or {}
        workout = extracted.get("workout")

        cal = overrides.get("calories_kcal", nutrition.get("calories_kcal"))
        pro = overrides.get("protein_g", nutrition.get("protein_g"))
        carbs = overrides.get("carbs_g", nutrition.get("carbs_g"))
        fat = overrides.get("fat_g", nutrition.get("fat_g"))
        water = overrides.get("hydration_ml", extracted.get("hydration_ml"))
        workout_str = workout.get("activity", "") if isinstance(workout, dict) else ""

        logged_at = log.get("logged_at") or log.get("timestamp")
        time_str = _format_dt(logged_at)

        row_data = [
            log.get("date", ""),
            time_str,
            log.get("text", ""),
            cal if cal is not None else "",
            pro if pro is not None else "",
            carbs if carbs is not None else "",
            fat if fat is not None else "",
            water if water is not None else "",
            workout_str,
            log.get("source", "text").capitalize(),
        ]

        ws_logs.row_dimensions[row_num].height = 20
        is_even = (row_num % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_logs.cell(row=row_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            if col_idx in (4, 5, 6, 7, 8):
                cell.alignment = data_align_right
            elif col_idx in (1, 2, 10):
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left

        row_num += 1

    apply_sheet_styling(ws_logs, logs_headers)

    # ── SHEET 2: Daily Summaries ──────────────────────────────────────────────
    ws_days = wb.create_sheet(title="Daily Summaries")
    days_headers = [
        "Date",
        "Total Calories (kcal)",
        "Calorie Target",
        "Total Protein (g)",
        "Protein Target",
        "Water Logged (ml)",
        "Water Target",
        "Workout Logged",
        "Status",
    ]
    row_num = 2
    for day in days:
        totals = day.get("totals") or {}
        cal_val = totals.get("calories", totals.get("calories_kcal", 0))
        pro_val = totals.get("protein_g", 0)
        water_val = totals.get("water_ml", 0)
        workout_done = "Yes" if day.get("workout_done") or totals.get("workout_done") else "No"
        status_str = (day.get("status") or "open").capitalize()

        row_data = [
            day.get("date", ""),
            cal_val,
            user.get("daily_calorie_target") or "",
            pro_val,
            user.get("daily_protein_target_g") or "",
            water_val,
            user.get("daily_water_target_ml") or 3000,
            workout_done,
            status_str,
        ]

        ws_days.row_dimensions[row_num].height = 20
        is_even = (row_num % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_days.cell(row=row_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            if col_idx in (2, 3, 4, 5, 6, 7):
                cell.alignment = data_align_right
            else:
                cell.alignment = data_align_center

        row_num += 1

    apply_sheet_styling(ws_days, days_headers)

    # ── SHEET 3: Saved Favorite Meals ─────────────────────────────────────────
    ws_meals = wb.create_sheet(title="Saved Favorite Meals")
    meals_headers = [
        "Meal Name",
        "Ingredients / Description",
        "Calories (kcal)",
        "Protein (g)",
        "Carbs (g)",
        "Fats (g)",
        "Times Logged",
        "Pinned",
    ]
    row_num = 2
    for meal in saved_meals:
        ext = meal.get("extracted") or {}
        cal = ext.get("calories") or ext.get("calories_kcal")
        pro = ext.get("protein_g")
        carbs = ext.get("carbs_g")
        fat = ext.get("fat_g")
        is_pinned = "Yes" if meal.get("pinned") else "No"

        row_data = [
            meal.get("name", ""),
            meal.get("raw_text", ""),
            cal if cal is not None else "",
            pro if pro is not None else "",
            carbs if carbs is not None else "",
            fat if fat is not None else "",
            meal.get("use_count", 0),
            is_pinned,
        ]

        ws_meals.row_dimensions[row_num].height = 20
        is_even = (row_num % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_meals.cell(row=row_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            if col_idx in (3, 4, 5, 6, 7):
                cell.alignment = data_align_right
            elif col_idx in (8,):
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left

        row_num += 1

    apply_sheet_styling(ws_meals, meals_headers)

    # ── SHEET 4: Profile & Goals ──────────────────────────────────────────────
    ws_profile = wb.create_sheet(title="Profile & Goals")
    profile_headers = ["Parameter / Target", "Value"]
    profile_rows = [
        ("User Name", user.get("name", "")),
        ("Username", user.get("username", "")),
        ("Primary Goal", (user.get("goal") or "").replace("_", " ").title()),
        ("Body Weight (kg)", user.get("body_weight_kg") or ""),
        ("Daily Calorie Target (kcal)", user.get("daily_calorie_target") or ""),
        ("Daily Protein Target (g)", user.get("daily_protein_target_g") or ""),
        ("Daily Water Target (ml)", user.get("daily_water_target_ml") or 3000),
        ("Weekly Workout Target (days)", user.get("weekly_workout_target") or 3),
        ("Export Generated At", datetime.now().strftime("%Y-%m-%d %I:%M %p")),
    ]
    row_num = 2
    for label, val in profile_rows:
        ws_profile.row_dimensions[row_num].height = 20
        c1 = ws_profile.cell(row=row_num, column=1, value=label)
        c2 = ws_profile.cell(row=row_num, column=2, value=val)
        c1.font = Font(name="Calibri", size=10, bold=True, color="1E293B")
        c2.font = data_font
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = data_align_left
        c2.alignment = data_align_left
        row_num += 1

    apply_sheet_styling(ws_profile, profile_headers)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@router.get("")
def export_user_data(
    format: str = Query(default="xlsx", pattern="^(json|csv|xlsx)$"),
    auth_context: dict = Depends(get_auth_context),
    database: Database = Depends(get_database),
):
    """
    Exports all data owned by the authenticated caller.
    format=xlsx (default): streams formatted Excel (.xlsx) workbook with human-readable sheets.
    format=json: streams structured JSON containing logs, days, sleep/wake, profile, saved meals.
    format=csv: streams CSV containing all food/hydration/workout logs.
    """
    user = auth_context["user"]
    user_id = user["user_id"]
    now = datetime.now(timezone.utc)
    date_slug = now.strftime("%Y-%m-%d")

    # Fetch user data
    logs_raw = list(
        database["daily_logs"]
        .find({"user_id": user_id})
        .sort([("date", 1), ("timestamp", 1)])
    )
    days_raw = list(database["days"].find({"user_id": user_id}).sort("date", 1))
    saved_meals_raw = list(database["saved_meals"].find({"user_id": user_id}))

    if format == "xlsx":
        excel_bytes = _build_excel_export(user, logs_raw, days_raw, saved_meals_raw)
        return Response(
            content=excel_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="ZenLog_Export_{date_slug}.xlsx"'
            },
        )

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "id",
            "date",
            "logged_at",
            "text",
            "protein_g",
            "calories_kcal",
            "hydration_ml",
            "workout",
            "source",
            "extraction_status",
        ])

        for log in logs_raw:
            extracted = log.get("extracted") or {}
            nutrition = extracted.get("nutrition") or {}
            overrides = log.get("overrides") or {}
            workout = extracted.get("workout")

            protein = overrides.get("protein_g", nutrition.get("protein_g", ""))
            calories = overrides.get("calories_kcal", nutrition.get("calories_kcal", ""))
            hydration = overrides.get("hydration_ml", extracted.get("hydration_ml", ""))
            workout_str = workout.get("activity", "") if isinstance(workout, dict) else ""

            logged_at_dt = log.get("logged_at") or log.get("timestamp")
            logged_at_str = logged_at_dt.isoformat() if isinstance(logged_at_dt, datetime) else str(logged_at_dt or "")

            writer.writerow([
                str(log.get("_id") or ""),
                log.get("date", ""),
                logged_at_str,
                log.get("text", ""),
                protein,
                calories,
                hydration,
                workout_str,
                log.get("source", "text"),
                log.get("extraction_status", ""),
            ])

        csv_data = output.getvalue()
        return Response(
            content=csv_data,
            media_type="text/csv",
            headers={
                "Content-Disposition": f'attachment; filename="zenlog_logs_{user_id}_{date_slug}.csv"'
            },
        )

    # JSON format
    daily_records_raw = list(database["daily_records"].find({"user_id": user_id}).sort("date", 1))

    export_payload = {
        "version": "1.0",
        "exported_at": now.isoformat(),
        "user": {
            "user_id": user.get("user_id"),
            "username": user.get("username"),
            "name": user.get("name"),
            "goal": user.get("goal"),
            "daily_protein_target_g": user.get("daily_protein_target_g"),
            "daily_calorie_target": user.get("daily_calorie_target"),
            "daily_water_target_ml": user.get("daily_water_target_ml"),
            "weekly_workout_target": user.get("weekly_workout_target"),
            "body_weight_kg": user.get("body_weight_kg"),
        },
        "logs": [_serialize_doc(l) for l in logs_raw],
        "days": [_serialize_doc(d) for d in days_raw],
        "sleep_wake_cycles": [_serialize_doc(r) for r in daily_records_raw],
        "saved_meals": [_serialize_doc(m) for m in saved_meals_raw],
    }

    return JSONResponse(
        content=export_payload,
        headers={
            "Content-Disposition": f'attachment; filename="zenlog_export_{user_id}_{date_slug}.json"'
        },
    )
